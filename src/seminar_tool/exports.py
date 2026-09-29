"""Excel exports, the upload template + validation (screen ד) and the SPSS package (screen ב)."""
from __future__ import annotations

import datetime as dt
import io
import re
from pathlib import Path
from typing import Any

import pandas as pd

from . import paths

BINARY = ["merits_appeal", "homicide_murder", "media_any", "media_district_any", "media_appeal_any", "media_primo_any",
          "media_combined_appeal_any", "intervention", "defendant_helped", "timeline_clean", "media_ref_any",
          "media_ref_coverage", "media_ref_pressure", "media_ref_reasoning", "attempt_only", "media_district_window"]
REQUIRED = ["media_any", "intervention"]
TEXT_COLS = {"case_id", "case_number", "decision_date", "non_merits_reason", "reasoning_method"}


# ---------------------------------------------------------------- SPSS syntax helpers
def spss_dir() -> Path:
    return paths.DATA_DIR / "spss"


def spss_variable_order() -> list[str]:
    """Column order expected by GET DATA in 01_import_and_labels.sps."""
    txt = (spss_dir() / "01_import_and_labels.sps").read_text(encoding="utf-8-sig")
    block = txt[txt.index("/VARIABLES=") + len("/VARIABLES="): txt.index("CACHE.")]
    return re.findall(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s+[AF]\d", block, flags=re.M)


def spss_variable_labels() -> dict[str, str]:
    txt = (spss_dir() / "01_import_and_labels.sps").read_text(encoding="utf-8-sig")
    block = txt[txt.index("VARIABLE LABELS"):]
    block = block[: block.index("'.") + 2]
    return dict(re.findall(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s+'([^']*)'", block, flags=re.M))


# ---------------------------------------------------------------- generic Excel writing
def _autosize(ws: Any) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)


def frames_to_excel(frames: list[tuple[str, pd.DataFrame]]) -> bytes:
    from openpyxl.styles import Font

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in frames:
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.sheet_view.rightToLeft = True
            for c in ws[1]:
                c.font = Font(bold=True)
            _autosize(ws)
    return buf.getvalue()


def stats_to_excel(result: dict[str, Any], golden: dict[str, Any] | None, source_label: str) -> bytes:
    """One sheet per analysis section (tables stacked with captions) + a golden-check sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True
    ws.append(["שחזור הניתוח הסטטיסטי — כלי נתוני הסמינריון"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["מקור הנתונים", source_label])
    ws.append(["נוצר ב-", dt.datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append(["מדגם הניתוח / קורפוס מלא", f"{result['n_main']} / {result['n_full']}"])
    if golden:
        ws.append(["השוואה לנתוני העבודה", f"{golden['matched']} מתוך {golden['checked']} ערכים זהים"
                   + (" ✓" if golden["all_ok"] else " — יש הבדלים, ראו גיליון 'השוואה'")])
    _autosize(ws)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    for sec in result["sections"]:
        sh = wb.create_sheet(sec["title"][:28].replace(":", "").replace("/", "-"))
        sh.sheet_view.rightToLeft = True
        sh.append([sec["title"]])
        sh.cell(sh.max_row, 1).font = Font(bold=True, size=13)
        sh.append([sec["help"]])
        for b in sec["blocks"]:
            sh.append([])
            if b["type"] == "note":
                sh.append([b["text"]])
                continue
            if b.get("caption"):
                sh.append([b["caption"]])
                sh.cell(sh.max_row, 1).font = Font(bold=True)
            sh.append(b["columns"])
            for c in sh[sh.max_row]:
                c.font, c.fill = Font(bold=True), head_fill
            for row in b["rows"]:
                sh.append(row)
        _autosize(sh)
        for row in sh.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=False, horizontal="right")
    if golden:
        sh = wb.create_sheet("השוואה")
        sh.sheet_view.rightToLeft = True
        sh.append(["ערך", "בעבודה (SPSS)", "חושב כעת", "זהה?"])
        for c in sh[1]:
            c.font, c.fill = Font(bold=True), head_fill
        for it in golden["items"]:
            sh.append([it["label"], it["expected"], it["got"], "✓" if it["ok"] else "✗"])
        _autosize(sh)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- template + upload (screen ד)
def template_excel(example: pd.DataFrame | None = None) -> bytes:
    labels = spss_variable_labels()
    cols = spss_variable_order()
    data = pd.DataFrame(columns=cols) if example is None else example[cols].head(3)
    rows = []
    for c in cols:
        kind = "0/1" if c in BINARY else ("טקסט" if c in TEXT_COLS else "מספר")
        need = "חובה" if c in REQUIRED else ("מומלץ" if c in ("merits_appeal", "year_c", "homicide_murder", "media_count") else "רשות")
        rows.append({"עמודה": c, "תיאור": labels.get(c, ""), "סוג ערך": kind, "חובה?": need})
    info = pd.DataFrame(rows)
    notes = pd.DataFrame({"הנחיות": [
        "ממלאים שורה אחת לכל תיק בגיליון 'נתונים'. את שמות העמודות בשורה הראשונה אין לשנות.",
        "עמודות חובה: media_any (בולטות 0/1) ו-intervention (התערבות 0/1). שאר העמודות מאפשרות ניתוחים נוספים.",
        "merits_appeal: 1 = נכלל במדגם הניתוח. אם העמודה חסרה — כל השורות נחשבות מדגם הניתוח.",
        "ערכי 0/1 בלבד בעמודות מסוג 0/1. תאים ריקים מותרים (השורה לא תיכלל בניתוח הרלוונטי).",
        "דוגמה מלאה: במסך 'ייבוא קובץ נתונים משלך' אפשר להוריד את נתוני המחקר עצמם בפורמט הזה.",
    ]})
    return frames_to_excel([("נתונים", data), ("הסבר העמודות", info), ("הנחיות", notes)])


def read_table(filename: str, raw: bytes) -> pd.DataFrame:
    """Read an uploaded Excel/CSV file; every failure becomes a Hebrew ValueError."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            xl = pd.ExcelFile(io.BytesIO(raw), engine="openpyxl")
            sheet = "נתונים" if "נתונים" in xl.sheet_names else xl.sheet_names[0]
            return xl.parse(sheet)
        except Exception as exc:  # noqa: BLE001
            raise ValueError("לא ניתן לקרוא את קובץ ה-Excel. ודאו שזה קובץ ‎.xlsx‏ תקין (אפשר לשמור אותו מחדש ב-Excel).") from exc
    if name.endswith(".xls"):
        raise ValueError("קובצי ‎.xls‏ ישנים אינם נתמכים. שמרו את הקובץ ב-Excel בפורמט ‎.xlsx‏ ונסו שוב.")
    if name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "cp1255"):
            try:
                return pd.read_csv(io.BytesIO(raw), encoding=enc)
            except UnicodeDecodeError:
                continue
            except Exception as exc:  # noqa: BLE001
                raise ValueError("לא ניתן לקרוא את קובץ ה-CSV. ודאו שהעמודות מופרדות בפסיקים ושהשורה הראשונה היא שמות העמודות.") from exc
        raise ValueError("קידוד התווים של קובץ ה-CSV אינו מוכר. שמרו אותו כ-CSV UTF-8 ונסו שוב.")
    raise ValueError("סוג הקובץ אינו נתמך. יש להעלות קובץ Excel ‏(‎.xlsx‏) או CSV.")


def validate_upload(df: pd.DataFrame) -> dict[str, Any]:
    """Hebrew error/warning messages; the DataFrame is cleaned in place (numeric columns)."""
    errors: list[str] = []
    warnings: list[str] = []
    df.columns = [str(c).strip() for c in df.columns]
    df.dropna(how="all", inplace=True)
    if df.empty:
        return {"ok": False, "errors": ["הקובץ ריק — לא נמצאו שורות נתונים."], "warnings": [], "n_rows": 0}
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        errors.append("חסרות עמודות חובה: " + ", ".join(missing)
                      + ". יש להשתמש בקובץ התבנית ולא לשנות את שמות העמודות בשורה הראשונה.")
    known = set(spss_variable_order())
    unknown = [c for c in df.columns if c not in known]
    if unknown:
        warnings.append("עמודות שאינן מוכרות ולא ישמשו בניתוח: " + ", ".join(unknown[:12]) + ("..." if len(unknown) > 12 else ""))
    for c in df.columns:
        if c in TEXT_COLS or c not in known:
            continue
        conv = pd.to_numeric(df[c], errors="coerce")
        bad = df[c].notna() & conv.isna()
        if bad.any():
            rows = ", ".join(str(i + 2) for i in df.index[bad][:8])
            errors.append(f"בעמודה {c} יש ערכים שאינם מספרים (שורות באקסל: {rows}{'...' if bad.sum() > 8 else ''}).")
        df[c] = conv
        if c in BINARY:
            bad01 = conv.notna() & ~conv.isin([0, 1])
            if bad01.any():
                rows = ", ".join(str(i + 2) for i in df.index[bad01][:8])
                errors.append(f"בעמודה {c} מותרים רק הערכים 0 או 1 (שורות באקסל: {rows}{'...' if bad01.sum() > 8 else ''}).")
    for c in REQUIRED:
        if c in df.columns and df[c].notna().sum() and df[c].dropna().nunique() < 2:
            errors.append(f"בעמודה {c} יש ערך אחד בלבד — אי אפשר להשוות בין קבוצות.")
    if "merits_appeal" not in df.columns:
        warnings.append("העמודה merits_appeal חסרה — כל השורות ייחשבו מדגם הניתוח.")
    optional_parts = {
        "רגרסיה מתוקננת": ["year_c", "homicide_murder"],
        "מינון (מספר ידיעות)": ["media_count"],
        "אזכורי תקשורת בפסק הדין": ["media_ref_any"],
        "מדד מורכבות טקסטואלית": ["r_sentence_len", "r_sentence_len_p90", "r_word_len", "r_long_words_pct",
                                  "r_very_long_words_pct", "r_paren_per_1000"],
    }
    skipped = [f"{k} (חסר: {', '.join(c for c in v if c not in df.columns)})" for k, v in optional_parts.items()
               if any(c not in df.columns for c in v)]
    if skipped:
        warnings.append("ניתוחים שידולגו בגלל עמודות חסרות: " + "; ".join(skipped))
    n_main = int((df["merits_appeal"] == 1).sum()) if "merits_appeal" in df.columns else len(df)
    return {"ok": not errors, "errors": errors, "warnings": warnings, "n_rows": int(len(df)), "n_main": n_main,
            "columns": list(df.columns)}


# ---------------------------------------------------------------- SPSS package
def _sps_rewrite(text: str, folder: Path) -> str:
    """Point the SPSS syntax at `folder` (templates carry a {{SPSS_FOLDER}} placeholder; originals a project path)."""
    root = folder.as_posix().replace("'", "''")
    text = text.replace("{{SPSS_FOLDER}}", root)
    text = re.sub(r"FILE HANDLE root /NAME='[^']*'\.", lambda m: f"FILE HANDLE root /NAME='{root}'.", text)
    return text.replace("root/02_נתונים/", "root/").replace("root/03_ניתוח/spss/output/", "root/output/")


def spss_export(folder: str | Path, df: pd.DataFrame | None = None) -> dict[str, Any]:
    """Write study_dataset_final.csv + the two .sps files (paths pointing to `folder`) + a Hebrew read-me."""
    folder = Path(folder).expanduser()
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "output").mkdir(exist_ok=True)
    order = spss_variable_order()
    target = folder / "study_dataset_final.csv"
    if df is None:
        target.write_bytes((paths.DATA_DIR / "study_dataset_final.csv").read_bytes())
    else:
        missing = [c for c in order if c not in df.columns]
        if missing:
            raise ValueError("לייצוא ל-SPSS נדרשות כל עמודות קובץ המחקר (לפי התבנית). חסרות: " + ", ".join(missing))
        df[order].to_csv(target, index=False, encoding="utf-8-sig")
    for name in ("01_import_and_labels.sps", "02_analysis.sps"):
        txt = (spss_dir() / name).read_text(encoding="utf-8-sig")
        (folder / name).write_text(_sps_rewrite(txt, folder), encoding="utf-8")  # same encoding as the originals
    readme = (
        "הרצת הניתוח ב-SPSS (גרסה 27 ומעלה)\n"
        "=====================================\n"
        f"התיקייה: {folder}\n\n"
        "1. פתחו את SPSS. בתפריט File > Open > Syntax בחרו את הקובץ 01_import_and_labels.sps.\n"
        "2. בחלון התחביר: Run > All. הקובץ קורא את study_dataset_final.csv, מוסיף תוויות ושומר study_dataset_final.sav.\n"
        "3. פתחו באותו אופן את 02_analysis.sps והריצו Run > All.\n"
        "   הפלט נשמר גם בתיקייה output (קובץ ‎.spv, טבלאות Excel ו-XML).\n\n"
        "הנתיבים בקבצי התחביר הותאמו לתיקייה הזו. אם מעבירים את התיקייה למקום אחר — יש לעדכן את השורה\n"
        "FILE HANDLE root /NAME='...' בראש שני הקבצים, או לייצא מחדש מהכלי.\n"
    )
    (folder / "קרא_אותי_SPSS.txt").write_text(readme, encoding="utf-8-sig")
    return {"folder": str(folder), "files": sorted(p.name for p in folder.iterdir())}
