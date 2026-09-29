"""Builds the Hebrew user guide 05_תוכנה\\מדריך_למשתמש.docx (+ .pdf through Microsoft Word, if installed).
    python src/guide/make_guide.py
Screenshots come from src/guide/screens (take_screenshots.py).
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SCREENS = HERE / "screens"
TOOL = HERE.parents[1]
OUT_DOCX = TOOL / "מדריך_למשתמש.docx"
OUT_PDF = TOOL / "מדריך_למשתמש.pdf"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
FONT = "Arial"


def rtl_paragraph(p) -> None:
    ppr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.insert(0, bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def rtl_run(run, size: float = 11, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rpr.append(rtl)
    cs = OxmlElement("w:szCs")
    cs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(cs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


class Guide:
    def __init__(self) -> None:
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sectpr = sec._sectPr
        bidi = OxmlElement("w:bidi")
        sectpr.append(bidi)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(3)

    def para(self, *parts, size: float = 11, space_after: float = 3, bullet: bool = False) -> None:
        """parts: strings, or (text, 'b') for bold."""
        p = self.doc.add_paragraph()
        rtl_paragraph(p)
        p.paragraph_format.space_after = Pt(space_after)
        if bullet:
            p.paragraph_format.right_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            r = p.add_run("•  ")
            rtl_run(r, size)
        for part in parts:
            text, bold = (part, False) if isinstance(part, str) else (part[0], True)
            if isinstance(part, tuple) and len(part) > 1 and part[1] == "ltr":
                text, bold = "‎" + part[0] + "‎", False
            rtl_run(p.add_run(text), size, bold)

    def heading(self, text: str, size: float = 15, space_before: float = 8) -> None:
        p = self.doc.add_paragraph()
        rtl_paragraph(p)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        rtl_run(p.add_run(text), size, True, NAVY)

    def image(self, name: str, width_cm: float = 17.0, caption: str = "") -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = bool(caption)
        p.add_run().add_picture(str(SCREENS / f"{name}.png"), width=Cm(width_cm))
        if caption:
            c = self.doc.add_paragraph()
            rtl_paragraph(c)
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after = Pt(6)
            rtl_run(c.add_run(caption), 9, False, RGBColor(0x5B, 0x67, 0x76))

    def table(self, rows: list[list[str]], widths: list[float]) -> None:
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        bidi = OxmlElement("w:bidiVisual")
        tblpr.append(bidi)
        for i, row in enumerate(rows):
            for j, text in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                p = cell.paragraphs[0]
                rtl_paragraph(p)
                rtl_run(p.add_run(text), 10, i == 0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self) -> None:
        self.doc.add_page_break()


def build() -> None:
    g = Guide()
    # ---------------------------------------------------------------- title
    p = g.doc.add_paragraph()
    rtl_paragraph(p)
    rtl_run(p.add_run("כלי נתוני הסמינריון — מדריך למשתמש"), 22, True, NAVY)
    g.para("בולטות תקשורתית מוקדמת והתערבות ערעורית: ערעורי רצח והמתה בבית המשפט העליון, 2010–2020", size=12)
    g.para(f"גרסה 1.1 · {dt.date.today().strftime('%d.%m.%Y')}", size=9, space_after=8)

    g.heading("מה הכלי עושה?")
    g.para("הכלי מאפשר לעיין בנתוני העבודה הסמינריונית, לשחזר בלחיצה אחת את כל הניתוחים הסטטיסטיים שלה ולוודא שהם זהים "
           "למספרים שבעבודה, לייבא מחדש את הנתונים מהמקורות, ולהריץ את אותם ניתוחים על קובץ נתונים משלכם. ",
           ("הכלי עובד בכל מחשב Windows 10 / 11"), " (64 סיביות), מכל תיקייה — גם מכונן USB או מתיקיית רשת — ",
           ("ללא התקנה"), " וללא Python או תוכנה נוספת. הכלי פועל בדפדפן, במחשב שלכם בלבד; חיבור לאינטרנט נדרש רק במסך „ייבוא מהמקורות”.")

    g.heading("הפעלה (בלי התקנה)")
    g.para("1. ", ("חלצו קודם את כל קובץ ה-ZIP"), " (לחיצה ימנית ← „חלץ הכל” / Extract All) לתיקייה כלשהי, למשל לשולחן העבודה. "
           "אם מפעילים את התוכנה מתוך ה-ZIP בלי לחלץ, תופיע ההודעה „יש לחלץ את כל תוכן קובץ ה-ZIP…” — חלצו ונסו שוב.")
    g.para("2. בתיקייה שחולצה ", ("לחצו פעמיים על „הפעלה.bat”"), " (או על SeminarDataTool.exe). אחרי כמה שניות נפתח הדפדפן (Edge, Chrome, Firefox — "
           "מה שמוגדר במחשב) עם מסך הבית. אם הדפדפן לא נפתח, מופיע חלון קטן עם הכתובת להקלדה בדפדפן.")
    g.para("3. בהפעלה הראשונה Windows עשוי להציג חלון כחול „Windows protected your PC”: לחצו ", ("More info"), " ואז ", ("Run anyway"),
           ". זה קורה בכל תוכנה שאינה חתומה דיגיטלית; התוכנה אינה מתקינה דבר ואינה משנה דבר במחשב.")
    g.para("4. ", ("סגירה:"), " לחצו על הכפתור „סגירה” בראש המסך, או פשוט סגרו את לשונית הדפדפן — התוכנה נסגרת מעצמה תוך כמה שניות.")

    g.heading("שמירת קבצים — בכל תיקייה שתבחרו")
    g.para("בכל ייצוא (Excel, CSV, SPSS, תבנית, יומן התוכנה, תיקיית ההורדה של מאגר פסקי הדין) נפתח ", ("חלון „שמירה בשם” של Windows"),
           " ואפשר לבחור כל תיקייה — גם כונן USB או תיקיית רשת. החלון נפתח בתיקייה האחרונה שבחרתם. אחרי השמירה מוצג הנתיב המלא של הקובץ "
           "וכפתור „הצגה בתיקייה”. ברירת המחדל היא התיקייה ", ("„תוצרים”"), " שנוצרת ליד התוכנה (שם נשמר גם יומן התוכנה, log.txt).")
    g.para(("אם אי אפשר לכתוב בתיקיית התוכנה"), " (למשל Program Files, כונן או תיקיית רשת לקריאה בלבד, או תיקייה ללא הרשאה), הכלי משתמש אוטומטית "
           "בתיקייה ", ("מסמכים\\SeminarDataTool\\תוצרים"), " (ואם אין תיקיית מסמכים — AppData\\Local\\SeminarDataTool\\תוצרים) ומציג על כך הודעה במסך הבית. "
           "נתוני המחקר עצמם נמצאים בתיקייה data ואינם משתנים.")
    g.image("01_home", 13.5, "מסך הבית: ארבע אפשרויות העבודה ומיקומי השמירה")

    # ---------------------------------------------------------------- A
    g.heading("א. צפייה בנתוני המחקר")
    g.para("טבלה של כל 497 התיקים. אפשר לסנן לפי שנה, סוג עבירה (רצח או ניסיון לרצח / המתה אחרת), בולטות תקשורתית (עם / בלי ידיעות), "
           "התערבות בית המשפט העליון ומדגם (478 הערעורים לגופו, או 19 ההליכים שהוחרגו), ולחפש מספר הליך. לחיצה על כותרת עמודה ממיינת; "
           "לחיצה על שורה פותחת את ", ("כרטיס התיק"), ".")
    g.para("כרטיס התיק מציג: פרטי ההליך והתוצאה, ", ("מדוע התיק נכלל"), " (או מדוע הוחרג), הידיעות שנמצאו ב-Google News "
           "(כותרת, כלי תקשורת, תאריך, קישור, חלון הזמן וסטטוס הניקוי הידני), ", ("השאילתות המדויקות"),
           " שהופעלו (עם קישור לאותו חיפוש), רשומות Primo, ואת ", ("טקסט פסק הדין"),
           " — בלחיצה על „טקסט פסק הדין” הוא נפתח, ואזכורי התקשורת שאומתו ידנית מסומנים בצבע לפי סוגם; קו מקווקו מסמן את תחילת פרק ההנמקה.")
    g.image("04_case_judgment", 13.5, "טקסט פסק הדין עם אזכור תקשורת מסומן")

    # ---------------------------------------------------------------- B
    g.heading("ב. שחזור הניתוח הסטטיסטי")
    g.para("המסך מחשב מחדש, מקובץ הנתונים, את כל הטבלאות של פרק הממצאים: תיאור המדגם; טבלת 2×2 עם χ², תיקון Yates, Fisher "
           "(דו-צדדי וחד-צדדי), φ, יחס סיכויים ורווח סמך; רגרסיה לוגיסטית (עם ובלי בקרה לשנה ולסוג העבירה); בדיקות חוסן; מינון; "
           "אזכורי תקשורת בפסק הדין; מדד המורכבות הטקסטואלית של ההנמקה; תיקון Holm למבחני השאלה המשנית; "
           "ועוצמה סטטיסטית (ניתוח רגישות חד-צדדי).")
    g.para("בראש המסך מופיע ", ("סימון ירוק „✓ זהה לנתוני העבודה”"), " — הכלי משווה כל מספר לגיליון המספרים המאושר של העבודה "
           "(פלט SPSS 27). ליד כל פרק מופיע סימון נפרד. „הצגת ההשוואה המלאה” מציגה את כל הערכים זה לצד זה.")
    g.image("05_stats_top", 13.5, "הסימון הירוק, ומתחת לכפתורים — היכן נשמר קובץ ה-Excel")
    g.para(("ייצוא ל-Excel:"), " קובץ עם גיליון לכל פרק וגיליון השוואה. ", ("ייצוא ל-SPSS:"),
           " בוחרים תיקייה („בחירת תיקייה…”) והכלי שומר בה את קובץ הנתונים ואת שני קובצי התחביר של העבודה, כשהנתיבים בתוכם כבר מותאמים לתיקייה. "
           "ב-SPSS פותחים (", ("File > Open > Syntax", "ltr"), ") את הקובץ ", ("01_import_and_labels.sps", "ltr"), " ומריצים (",
           ("Run > All", "ltr"), "), ואחר כך באותו אופן את ", ("02_analysis.sps", "ltr"), ". הפלט נשמר גם בתת-התיקייה output.")

    # ---------------------------------------------------------------- C
    g.heading("ג. ייבוא הנתונים מחדש מהמקורות (דורש אינטרנט)")
    g.para(("שלב 1 — נתוני בית המשפט."), " פסקי הדין נלקחים ממאגר הנתונים הפתוח LevMuchnik/SupremeCourtOfIsrael באתר Hugging Face "
           "(מאגר נתונים — לא מודל בינה מלאכותית; קובץ אחד של כ-1.5 GB). אם כבר יש עותק במחשב — הכלי מאתר אותו; אחרת לוחצים "
           "„הורדת הקובץ” (5–30 דקות; אפשר לעצור ולהמשיך; את תיקיית ההורדה אפשר לשנות). „בניית האוכלוסייה” מריצה את אותם סינונים "
           "של העבודה ומציגה את מספר התיקים בכל שלב לצד המספר שבעבודה.")
    g.image("09_reimport_flow", 13.5, "שלבי הסינון: המספרים בשחזור לעומת העבודה")
    g.para(("שלב 2 — ידיעות Google News."), " בוחרים שנה או תיקים בודדים (עד 60 בכל ריצה), והכלי מפעיל את אותה שאילתת „שם מדויק” "
           "בשני חלונות הזמן, בקצב איטי ומנומס (השהיה של כמה שניות בין בקשות; עצירה אוטומטית אם Google מבקש להאט; כפתור „עצירה”; "
           "אפשר להמשיך מאותה נקודה). נשמרות רק ידיעות שתאריכן בתוך החלון, והן מוצגות לצד הידיעות שנכללו במחקר. ",
           ("חשוב:"), " תוצאות Google News משתנות עם הזמן והרשימה החדשה לא עברה סינון ידני — ", ("תמונת המצב המקורית של המחקר היא נקודת הייחוס"), ".")
    g.para(("Primo:"), " החיפוש במאגר העיתונות של הספרייה דורש כניסה אישית, ולכן אינו מבוצע בכלי. הרשומות שנאספו במחקר מוצגות בכרטיס התיק.")
    g.para(("שלב 3 — ייצוא:"), " שמירת האוכלוסייה המשוחזרת ותוצאות החיפוש ל-Excel או ל-CSV, בתיקייה שתבחרו.")

    # ---------------------------------------------------------------- D
    g.heading("ד. ייבוא קובץ נתונים משלך")
    g.para("1. שומרים את ", ("תבנית ה-Excel"), " (גיליון „הסבר העמודות” מתאר כל עמודה). אפשר גם לשמור את נתוני המחקר עצמם כדוגמה מלאה.")
    g.para("2. ממלאים שורה לכל תיק. עמודות חובה: media_any (בולטות 0/1) ו-intervention (התערבות 0/1); עמודות נוספות מאפשרות ניתוחים נוספים.")
    g.para("3. מעלים את הקובץ (Excel או CSV). הכלי בודק אותו ומציג הודעות ברורות: עמודות חסרות, ערכים שאינם 0/1 (עם מספרי השורות באקסל) וכו'.")
    g.para("4. „הרצת הניתוחים על הקובץ” מציגה את אותם ניתוחים של מסך ב על הנתונים שלכם, עם ייצוא ל-Excel ול-SPSS. הקובץ לא יוצא מהמחשב.")

    # ---------------------------------------------------------------- FAQ
    g.heading("תקלות נפוצות")
    g.table([
        ["מה קורה", "מה עושים"],
        ["„יש לחלץ את כל תוכן קובץ ה-ZIP…”", "התוכנה הופעלה מתוך ה-ZIP. לחיצה ימנית על ה-ZIP ← „חלץ הכל”, ואז הפעלה מהתיקייה שחולצה."],
        ["הדפדפן לא נפתח", "מופיע חלון עם הכתובת (למשל ‎http://127.0.0.1:8765/‎) — פתחו דפדפן והקלידו אותה."],
        ["„Windows protected your PC”", "לחצו More info ואז Run anyway (פעם אחת בלבד)."],
        ["אנטי-וירוס חוסם את התוכנה", "התרעת שווא נפוצה בתוכנות לא חתומות; אפשרו את התיקייה SeminarDataTool."],
        ["הודעה שהקבצים יישמרו ב„מסמכים”", "תיקיית התוכנה מוגנת מפני כתיבה; זה תקין — הקבצים נשמרים במסמכים\\SeminarDataTool\\תוצרים."],
        ["הודעה „התוכנה אינה פועלת”", "התוכנה נסגרה (למשל אחרי שהמחשב נכנס למצב שינה). הפעילו שוב את „הפעלה.bat”."],
        ["הודעת שגיאה", "פרטים נשמרים בקובץ log.txt (בתיקיית „תוצרים”); במסך הבית: „שמירת עותק של יומן התוכנה”."],
    ], [5.0, 12.0])

    cp = g.doc.core_properties                    # no user or machine names in the document properties
    cp.author = cp.last_modified_by = "SeminarDataTool"
    cp.title = "כלי נתוני הסמינריון — מדריך למשתמש"
    cp.comments = ""
    g.doc.save(OUT_DOCX)
    print("saved", OUT_DOCX)


def to_pdf() -> None:
    try:
        import win32com.client  # type: ignore
    except ImportError:
        print("pywin32 not installed - PDF skipped")
        return
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(OUT_DOCX), ReadOnly=True)
        doc.SaveAs2(str(OUT_PDF), FileFormat=17)
        pages = doc.ComputeStatistics(2)
        doc.Close(False)
        print("saved", OUT_PDF, "pages:", pages)
    finally:
        word.Quit()


if __name__ == "__main__":
    build()
    to_pdf()
